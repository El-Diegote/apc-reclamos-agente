"""Carga versionada del Excel RESOLUCIONES APC."""

from __future__ import annotations

import hashlib
import logging
import re
import unicodedata
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.utils import parsear_importe


@dataclass(slots=True)
class ResolutionCatalogEntry:
    """Resolucion APC extraida de una fila del catalogo."""

    tema: str
    nota: str
    canal: str
    hoja: str
    fila: int
    origen_archivo: str
    version_origen: str
    vigencia_desde: str
    vigencia_hasta: str
    activa: bool = True


@dataclass(slots=True)
class DollarCoefficientEntry:
    """Coeficiente dolar extraido del catalogo."""

    descripcion: str
    valor: float
    hoja: str
    fila: int
    origen_archivo: str
    version_origen: str


class ResolutionCatalogLoader:
    """Lee RESOLUCIONES APC sin asumir posiciones rigidas."""

    HEADER_SCAN_ROWS = 20
    COEFFICIENT_SCAN_WINDOW_ROWS = 25
    COEFFICIENT_SCAN_WINDOW_COLS = 12

    def __init__(self, logger: logging.Logger | None = None) -> None:
        """Inicializa el cargador.

        Args:
            logger: Logger opcional.
        """
        self.logger = logger or logging.getLogger(self.__class__.__name__)

    def cargar(self, archivo: Path) -> dict[str, Any]:
        """Carga resoluciones y coeficientes desde un Excel local.

        Args:
            archivo: Ruta del Excel RESOLUCIONES APC.

        Returns:
            Catalogo serializable con metadata, resoluciones y coeficientes.
        """
        if not archivo.exists():
            raise FileNotFoundError(f"No existe el archivo de resoluciones: {archivo}")

        workbook = load_workbook(archivo, data_only=True)
        version = self._version_archivo(archivo)
        resoluciones: list[ResolutionCatalogEntry] = []
        coeficientes: list[DollarCoefficientEntry] = []

        for hoja in workbook.worksheets:
            header = self._detectar_header_resoluciones(hoja)
            if header:
                resoluciones.extend(self._leer_resoluciones(hoja, header, archivo.name, version))
            coeficientes.extend(self._leer_coeficientes_dolar(hoja, archivo.name, version))

        catalogo = {
            "metadata": {
                "origen_archivo": archivo.name,
                "version_origen": version,
                "hojas": workbook.sheetnames,
                "total_resoluciones": len(resoluciones),
                "total_coeficientes_dolar": len(coeficientes),
            },
            "resoluciones": [asdict(resolucion) for resolucion in resoluciones],
            "coeficientes_dolar": [asdict(coeficiente) for coeficiente in coeficientes],
        }
        self.logger.info(
            "Catalogo de resoluciones cargado: %s resoluciones, %s coeficientes",
            len(resoluciones),
            len(coeficientes),
        )
        return catalogo

    def seleccionar_coeficiente_dolar(self, catalogo: dict[str, Any]) -> float | None:
        """Selecciona el ultimo coeficiente dolar disponible en el catalogo.

        Args:
            catalogo: Resultado de `cargar`.

        Returns:
            Coeficiente encontrado o `None`.
        """
        coeficientes = catalogo.get("coeficientes_dolar", []) or []
        if not coeficientes:
            return None
        return float(coeficientes[-1]["valor"])

    def _detectar_header_resoluciones(self, hoja: Worksheet) -> dict[str, int] | None:
        """Detecta la fila de encabezados TEMA/NOTA."""
        for fila in hoja.iter_rows(max_row=min(hoja.max_row, self.HEADER_SCAN_ROWS)):
            columnas: dict[str, int] = {}
            for celda in fila:
                texto = self._normalizar(celda.value)
                if texto == "tema":
                    columnas["tema"] = celda.column
                elif texto in {"nota", "detalle", "resolucion", "resolucion sugerida"}:
                    columnas["nota"] = celda.column
                elif texto in {"vigencia desde", "desde", "fecha desde"}:
                    columnas["vigencia_desde"] = celda.column
                elif texto in {"vigencia hasta", "hasta", "fecha hasta"}:
                    columnas["vigencia_hasta"] = celda.column
            if "tema" in columnas and "nota" in columnas:
                columnas["header_row"] = fila[0].row
                return columnas
        return None

    def _leer_resoluciones(
        self,
        hoja: Worksheet,
        header: dict[str, int],
        origen_archivo: str,
        version_origen: str,
    ) -> list[ResolutionCatalogEntry]:
        """Lee resoluciones desde una hoja con encabezados detectados."""
        resoluciones: list[ResolutionCatalogEntry] = []
        for row_index in range(header["header_row"] + 1, hoja.max_row + 1):
            tema = self._texto(hoja.cell(row=row_index, column=header["tema"]).value)
            nota = self._texto(hoja.cell(row=row_index, column=header["nota"]).value)
            if not tema or not nota:
                continue
            resoluciones.append(
                ResolutionCatalogEntry(
                    tema=tema,
                    nota=nota,
                    canal=self._inferir_canal(tema),
                    hoja=hoja.title,
                    fila=row_index,
                    origen_archivo=origen_archivo,
                    version_origen=version_origen,
                    vigencia_desde=self._texto_celda_opcional(hoja, row_index, header.get("vigencia_desde")),
                    vigencia_hasta=self._texto_celda_opcional(hoja, row_index, header.get("vigencia_hasta")),
                )
            )
        return resoluciones

    def _leer_coeficientes_dolar(
        self,
        hoja: Worksheet,
        origen_archivo: str,
        version_origen: str,
    ) -> list[DollarCoefficientEntry]:
        """Busca zonas de coeficiente dolar y extrae valores numericos cercanos."""
        coeficientes: list[DollarCoefficientEntry] = []
        zonas = []
        for fila in hoja.iter_rows():
            for celda in fila:
                texto = self._normalizar(celda.value)
                if "coeficiente" in texto and "dolar" in texto:
                    zonas.append((celda.row, celda.column))

        for row_base, col_base in zonas:
            for row_index in range(
                row_base + 1,
                min(hoja.max_row, row_base + self.COEFFICIENT_SCAN_WINDOW_ROWS) + 1,
            ):
                descripcion = self._descripcion_fila(hoja, row_index, col_base)
                for col_index in range(
                    col_base,
                    min(hoja.max_column, col_base + self.COEFFICIENT_SCAN_WINDOW_COLS) + 1,
                ):
                    valor = self._parsear_coeficiente(
                        hoja.cell(row=row_index, column=col_index).value
                    )
                    if valor <= 0:
                        continue
                    coeficientes.append(
                        DollarCoefficientEntry(
                            descripcion=descripcion,
                            valor=valor,
                            hoja=hoja.title,
                            fila=row_index,
                            origen_archivo=origen_archivo,
                            version_origen=version_origen,
                        )
                    )
        return coeficientes

    def _descripcion_fila(self, hoja: Worksheet, row_index: int, col_base: int) -> str:
        """Obtiene la descripcion textual mas cercana a un coeficiente."""
        textos: list[str] = []
        for col_index in range(col_base, min(hoja.max_column, col_base + 5) + 1):
            valor = hoja.cell(row=row_index, column=col_index).value
            texto = self._texto(valor)
            if texto and parsear_importe(texto) == 0.0:
                textos.append(texto)
        return " | ".join(textos)

    def _parsear_coeficiente(self, valor: Any) -> float:
        """Convierte solo valores numericos o texto numerico a coeficiente."""
        if isinstance(valor, datetime):
            return 0.0
        if isinstance(valor, (int, float)):
            return float(valor)
        if isinstance(valor, str):
            return parsear_importe(valor)
        return 0.0

    def _texto_celda_opcional(self, hoja: Worksheet, row_index: int, col_index: int | None) -> str:
        """Lee texto de una columna opcional."""
        if col_index is None:
            return ""
        return self._texto(hoja.cell(row=row_index, column=col_index).value)

    def _texto(self, valor: Any) -> str:
        """Convierte valores de Excel a texto estable."""
        if valor is None:
            return ""
        if isinstance(valor, datetime):
            return valor.date().isoformat()
        return str(valor).strip()

    def _normalizar(self, valor: Any) -> str:
        """Normaliza texto para comparar encabezados y senales."""
        texto = unicodedata.normalize("NFKD", self._texto(valor))
        texto = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
        texto = texto.replace("�", "o")
        texto = re.sub(r"[^a-zA-Z0-9]+", " ", texto).strip().lower()
        return re.sub(r"\s+", " ", texto)

    def _version_archivo(self, archivo: Path) -> str:
        """Calcula una version estable basada en hash y fecha de modificacion."""
        digest = hashlib.sha256(archivo.read_bytes()).hexdigest()[:12]
        modificado = datetime.fromtimestamp(archivo.stat().st_mtime).strftime("%Y%m%d%H%M%S")
        return f"{modificado}-{digest}"

    def _inferir_canal(self, tema: str) -> str:
        """Infiere canal a partir del tema del catalogo."""
        tema_upper = tema.upper()
        if "ATM" in tema_upper:
            return "ATM"
        if "MPOS" in tema_upper:
            return "mPOS"
        if "POS" in tema_upper:
            return "POS"
        if "BNA" in tema_upper or "BILLETERA" in tema_upper:
            return "BNA+"
        if "MODO" in tema_upper:
            return "MODO"
        if "CASH" in tema_upper:
            return "Cash In"
        if "ECOMMERCE" in tema_upper or "E-COMMERCE" in tema_upper or "LINKGO" in tema_upper:
            return "eCommerce"
        return "Otros"
