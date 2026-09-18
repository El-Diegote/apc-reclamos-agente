"""Motor de negocio para el MVP APC Reclamos."""

from __future__ import annotations

import json
import logging
import re
import unicodedata
from dataclasses import asdict
from pathlib import Path
from typing import Any

import pandas as pd

from config.config import BASE_RECLAMOS_EXTENSIONS, DOCUMENT_EXTENSIONS, EXPORTS_DIR
from src.claim_universe_builder import ClaimUniverseBuilder
from src.database import RepositorioAPC
from src.decision_engine import OperationDecisionEngine
from src.document_classifier import DocumentClassifier
from src.resolution_engine import ResolutionEngine
from src.resolution_catalog import ResolutionCatalogLoader
from src.smart_console_analyzer import SmartConsoleAnalyzer
from src.utils import parsear_importe, timestamp_actual


class MotorAPC:
    """Coordina carga, entrenamiento, analisis y exportacion de casos APC."""

    def __init__(
        self,
        repositorio: RepositorioAPC | None = None,
        logger: logging.Logger | None = None,
    ) -> None:
        """Inicializa el motor.

        Args:
            repositorio: Repositorio SQLite opcional.
            logger: Logger opcional.
        """
        self.logger = logger or logging.getLogger(self.__class__.__name__)
        self.repositorio = repositorio or RepositorioAPC(logger=self.logger)
        self.coeficiente_dolar = 1.0

    def cargar_base_reclamos(self, archivo: Path) -> int:
        """Carga una base de reclamos desde CSV o Excel.

        Args:
            archivo: Ruta del archivo de base.

        Returns:
            Cantidad de registros persistidos.

        Raises:
            ValueError: Si el formato no esta soportado.
        """
        if archivo.suffix.lower() not in BASE_RECLAMOS_EXTENSIONS:
            raise ValueError(f"Formato no soportado para base de reclamos: {archivo.suffix}")

        if archivo.suffix.lower() == ".csv":
            dataframe = pd.read_csv(archivo)
        else:
            dataframe = pd.read_excel(archivo)

        filas = self.normalizar_base_reclamos(dataframe)
        cantidad = self.repositorio.guardar_base_reclamos(filas, archivo.name)
        self.logger.info("Base cargada desde %s", archivo)
        return cantidad

    def entrenar_resolucion(
        self,
        canal: str,
        motivo: str,
        texto_resolucion: str,
        texto_diario: str,
    ) -> int:
        """Entrena una resolucion APC reutilizable.

        Args:
            canal: Canal APC.
            motivo: Motivo del reclamo o patron textual.
            texto_resolucion: Resolucion sugerida.
            texto_diario: Texto diario sugerido.

        Returns:
            ID de la resolucion creada.
        """
        return self.repositorio.entrenar_resolucion(
            canal=canal.strip(),
            motivo=motivo.strip(),
            texto_resolucion=texto_resolucion.strip(),
            texto_diario=texto_diario.strip(),
        )

    def entrenar_desde_excel_resoluciones(self, archivo: Path) -> int:
        """Entrena resoluciones desde un Excel local con columnas TEMA y NOTA.

        Args:
            archivo: Ruta local del archivo de resoluciones APC.

        Returns:
            Cantidad de resoluciones entrenadas.

        Raises:
            ValueError: Si el archivo no contiene hojas aprovechables.
        """
        cargador = ResolutionCatalogLoader(logger=self.logger)
        catalogo = cargador.cargar(archivo)
        coeficiente = cargador.seleccionar_coeficiente_dolar(catalogo)
        if coeficiente:
            self.coeficiente_dolar = coeficiente
            self.logger.info("Coeficiente dolar actualizado desde catalogo de resoluciones")

        resoluciones_catalogo = catalogo["resoluciones"]
        if not resoluciones_catalogo:
            raise ValueError("No se encontraron resoluciones con columnas TEMA y NOTA.")

        total_inicial = len(self.repositorio.listar_resoluciones())
        for resolucion in catalogo["resoluciones"]:
            self.repositorio.entrenar_resolucion(
                canal=resolucion["canal"],
                motivo=resolucion["tema"],
                texto_resolucion=resolucion["nota"],
                texto_diario=self._generar_diario_desde_nota(
                    resolucion["tema"],
                    resolucion["nota"],
                ),
                origen_archivo=resolucion["origen_archivo"],
                version_origen=resolucion["version_origen"],
                hoja=resolucion["hoja"],
                fila=int(resolucion["fila"]),
                vigencia_desde=resolucion["vigencia_desde"],
                vigencia_hasta=resolucion["vigencia_hasta"],
            )

        entrenadas = len(self.repositorio.listar_resoluciones()) - total_inicial

        self.logger.info("Resoluciones entrenadas desde Excel: %s", entrenadas)
        return entrenadas

    def estado_resoluciones_apc(self, archivo: Path) -> dict[str, Any]:
        """Consulta si el Excel de resoluciones ya esta entrenado.

        Args:
            archivo: Ruta local del Excel RESOLUCIONES APC.

        Returns:
            Estado de version y carga.
        """
        cargador = ResolutionCatalogLoader(logger=self.logger)
        version = cargador.version_archivo(archivo)
        return {
            "archivo": str(archivo),
            "version_origen": version,
            "ya_cargada": self.repositorio.version_resoluciones_cargada(version),
        }

    def actualizar_resoluciones_apc(self, archivo: Path) -> dict[str, Any]:
        """Actualiza resoluciones solo si la version no fue cargada.

        Args:
            archivo: Ruta local del Excel RESOLUCIONES APC.

        Returns:
            Resultado de actualizacion con conteos y version.
        """
        estado = self.estado_resoluciones_apc(archivo)
        if estado["ya_cargada"]:
            return {
                **estado,
                "resoluciones_nuevas": 0,
                "actualizado": False,
                "mensaje": "La version vigente ya estaba entrenada.",
            }

        nuevas = self.entrenar_desde_excel_resoluciones(archivo)
        return {
            **estado,
            "resoluciones_nuevas": nuevas,
            "actualizado": nuevas > 0,
            "mensaje": f"Resoluciones APC actualizadas: {nuevas} nuevas.",
        }

    def analizar_caso(self, caso: dict[str, Any]) -> dict[str, Any]:
        """Analiza un caso y genera sugerencias revisables.

        Args:
            caso: Datos ingresados por el analista.

        Returns:
            Caso enriquecido con resolucion sugerida y texto diario.
        """
        caso_normalizado = self._normalizar_registro(caso)
        resolucion = self._buscar_mejor_resolucion(caso_normalizado)

        if resolucion:
            resolucion_sugerida = resolucion["texto_resolucion"]
            texto_diario = resolucion["texto_diario"]
            tema_diario = str(resolucion.get("motivo", "") or caso_normalizado.get("canal", ""))
            detalle_diario = str(resolucion.get("texto_resolucion", "") or "")
            fuente = f"resolucion_entrenada:{resolucion['id']}"
        else:
            resolucion_sugerida = self.generar_resolucion_sugerida(caso_normalizado)
            texto_diario = self.generar_texto_diario(caso_normalizado, resolucion_sugerida)
            tema_diario = str(caso_normalizado.get("canal", "Otros"))
            detalle_diario = resolucion_sugerida
            fuente = "regla_base_mvp"

        resultado = {
            **caso_normalizado,
            "resolucion_sugerida": resolucion_sugerida,
            "texto_diario": texto_diario,
            "tema_diario": tema_diario,
            "detalle_diario": detalle_diario,
            "coeficiente_dolar": self.coeficiente_dolar,
            "fuente_sugerencia": fuente,
            "requiere_aprobacion_humana": True,
            "aprobado": False,
            "timestamp_analisis": timestamp_actual(),
        }
        self.logger.info("Caso analizado: %s", resultado.get("numero_incidente"))
        return resultado

    def analizar_expediente(
        self,
        caso: dict[str, Any],
        documentos_dir: Path | None = None,
        archivos_extra: list[Path] | None = None,
    ) -> dict[str, Any]:
        """Analiza un expediente APC completo con los motores operativos.

        Args:
            caso: Datos visibles del caso actual.
            documentos_dir: Carpeta local con documentacion o exportaciones descargadas.
            archivos_extra: Archivos locales seleccionados manualmente.

        Returns:
            Analisis enriquecido con documentacion, Smart Console, decision y resolucion.
        """
        caso_normalizado = self._normalizar_registro(caso)
        archivos = self._listar_archivos_expediente(documentos_dir, archivos_extra)
        documentos = self._clasificar_documentos(archivos)
        smart_console = self._analizar_smart_console(archivos)

        universo = ClaimUniverseBuilder(logger=self.logger).construir(
            detalle_apc=caso_normalizado,
            diario_apc=[],
            documentacion=documentos,
            smart_console=smart_console,
        )
        decision_engine = OperationDecisionEngine(
            coeficiente_dolar=self.coeficiente_dolar,
            logger=self.logger,
        )
        resolucion_engine = ResolutionEngine(decision_engine=decision_engine, logger=self.logger)
        resolucion = resolucion_engine.generar(
            universo,
            resoluciones_entrenadas=self.repositorio.listar_resoluciones(),
        )

        analisis_base = self.analizar_caso(caso_normalizado)
        resultado = {
            **analisis_base,
            "documentos_clasificados": documentos,
            "smart_console": smart_console,
            "universo_reclamo": universo,
            "decision_operativa": resolucion["decision"],
            "evaluacion_fraude": resolucion["evaluacion_fraude"],
            "tema_diario": resolucion["tema"],
            "detalle_diario": resolucion["detalle"],
            "fuente_sugerencia": resolucion["fuente"],
            "requiere_aprobacion_humana": True,
        }
        self.logger.info("Expediente analizado: %s", resultado.get("numero_incidente"))
        return resultado

    def generar_resolucion_sugerida(self, caso: dict[str, Any]) -> str:
        """Genera una resolucion base cuando no hay entrenamiento especifico.

        Args:
            caso: Datos del caso.

        Returns:
            Texto de resolucion sugerida.
        """
        return (
            f"Se analizo el reclamo {caso.get('numero_incidente', '')} "
            f"correspondiente al canal {caso.get('canal', '')}. "
            f"Motivo informado: {caso.get('motivo', '')}. "
            "La resolucion queda sujeta a revision y aprobacion del analista APC."
        )

    def generar_texto_diario(self, caso: dict[str, Any], resolucion: str) -> str:
        """Genera texto sugerido para la solapa Diario.

        Args:
            caso: Datos del caso.
            resolucion: Resolucion sugerida.

        Returns:
            Texto diario sugerido.
        """
        importe = float(caso.get("importe") or 0)
        return (
            f"Analisis APC - Incidente {caso.get('numero_incidente', '')}. "
            f"Afectado: {caso.get('nombre_apellido', '')}. "
            f"Canal: {caso.get('canal', '')}. "
            f"Moneda: {caso.get('moneda', 'Peso')}. "
            f"Cuenta: {caso.get('numero_cuenta', '')}. "
            f"Importe reclamado: {importe:.2f}. "
            f"Resultado sugerido: {resolucion}"
        )

    def guardar_caso_aprobado(
        self,
        caso: dict[str, Any],
        aprobado: bool,
        analista: str,
    ) -> int:
        """Guarda un caso luego de la revision humana.

        Args:
            caso: Caso analizado.
            aprobado: Indica si el analista aprobo la sugerencia.
            analista: Nombre o identificador del analista.

        Returns:
            ID del caso guardado.
        """
        caso_guardar = {**caso, "aprobado": aprobado, "analista": analista}
        return self.repositorio.guardar_caso(caso_guardar)

    def exportar_informe(self, destino: Path | None = None) -> Path:
        """Exporta informe JSON de casos analizados.

        Args:
            destino: Ruta opcional del archivo de salida.

        Returns:
            Ruta del informe generado.
        """
        EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        destino = destino or EXPORTS_DIR / f"informe_apc_{timestamp_actual().replace(':', '-')}.json"
        casos = self.repositorio.listar_casos()
        informe = {
            "generado_en": timestamp_actual(),
            "total_casos": len(casos),
            "casos_aprobados": sum(1 for caso in casos if caso.get("aprobado")),
            "principio_operativo": (
                "Toda resolucion debe ser revisada y aprobada por un analista humano "
                "antes de publicarse."
            ),
            "casos": casos,
        }
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text(json.dumps(informe, ensure_ascii=False, indent=2), encoding="utf-8")
        self.logger.info("Informe exportado: %s", destino)
        return destino

    def exportar_texto_diario(self, texto_diario: str, formato: str = "txt") -> Path:
        """Exporta el texto Diario editado por el analista.

        Args:
            texto_diario: Texto final revisado.
            formato: Formato de salida, `txt` o `docx`.

        Returns:
            Ruta del archivo exportado.

        Raises:
            ValueError: Si el formato no esta soportado.
        """
        EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        timestamp = timestamp_actual().replace(":", "-")

        if formato == "txt":
            destino = EXPORTS_DIR / f"texto_diario_{timestamp}.txt"
            destino.write_text(texto_diario, encoding="utf-8")
            return destino

        if formato == "docx":
            from docx import Document

            destino = EXPORTS_DIR / f"texto_diario_{timestamp}.docx"
            documento = Document()
            documento.add_heading("Texto Diario APC", level=1)
            documento.add_paragraph(texto_diario)
            documento.save(destino)
            return destino

        raise ValueError(f"Formato no soportado: {formato}")

    def listar_resoluciones(self) -> list[dict[str, Any]]:
        """Lista resoluciones entrenadas.

        Returns:
            Resoluciones disponibles.
        """
        return self.repositorio.listar_resoluciones()

    def _listar_archivos_expediente(
        self,
        documentos_dir: Path | None,
        archivos_extra: list[Path] | None = None,
    ) -> list[Path]:
        """Lista archivos locales aptos para el analisis del expediente.

        Args:
            documentos_dir: Carpeta local de documentacion.
            archivos_extra: Archivos locales seleccionados manualmente.

        Returns:
            Archivos soportados encontrados.
        """
        archivos: list[Path] = []
        if documentos_dir is not None and documentos_dir.exists():
            archivos.extend(
                archivo
                for archivo in sorted(documentos_dir.iterdir())
                if archivo.is_file() and archivo.suffix.lower() in DOCUMENT_EXTENSIONS
            )
        for archivo in archivos_extra or []:
            if archivo.is_file() and archivo.suffix.lower() in DOCUMENT_EXTENSIONS:
                archivos.append(archivo)
        return list(dict.fromkeys(archivos))

    def _clasificar_documentos(self, archivos: list[Path]) -> list[dict[str, Any]]:
        """Clasifica documentos locales descargados o aportados.

        Args:
            archivos: Archivos del expediente.

        Returns:
            Clasificaciones documentales serializables.
        """
        clasificador = DocumentClassifier(logger=self.logger)
        clasificaciones: list[dict[str, Any]] = []
        for archivo in archivos:
            try:
                clasificaciones.append(asdict(clasificador.clasificar_archivo(archivo)))
            except Exception as exc:
                self.logger.warning("No se pudo clasificar %s: %s", archivo.name, exc)
                clasificaciones.append(
                    {
                        "archivo": archivo.name,
                        "tipo_documento": "Otro",
                        "confianza": 0.0,
                        "senales": [],
                        "requiere_ocr": archivo.suffix.lower() == ".pdf",
                        "error": str(exc),
                    }
                )
        return clasificaciones

    def _analizar_smart_console(self, archivos: list[Path]) -> dict[str, Any]:
        """Detecta y analiza la primera exportacion compatible de Smart Console.

        Args:
            archivos: Archivos disponibles del expediente.

        Returns:
            Resumen Smart Console o diccionario vacio si no hay archivo compatible.
        """
        analizador = SmartConsoleAnalyzer(logger=self.logger)
        for archivo in archivos:
            if archivo.suffix.lower() not in {".csv", ".xlsx", ".xlsm", ".xls"}:
                continue
            try:
                resumen = analizador.analizar_archivo(archivo)
                if resumen.get("total_operaciones", 0) > 0:
                    return resumen
            except Exception as exc:
                self.logger.info("Archivo omitido como Smart Console %s: %s", archivo.name, exc)
        return {}

    def normalizar_registro_publico(self, fila: dict[str, Any]) -> dict[str, Any]:
        """Normaliza un registro para usarlo desde la interfaz.

        Args:
            fila: Registro original.

        Returns:
            Registro normalizado.
        """
        return self._normalizar_registro(fila)

    def normalizar_base_reclamos(self, dataframe: pd.DataFrame) -> list[dict[str, Any]]:
        """Normaliza una base y consolida transacciones marcadas por incidente.

        Args:
            dataframe: Base CSV/Excel cargada localmente.

        Returns:
            Lista de casos normalizados para mostrar o persistir.
        """
        filas_originales = dataframe.to_dict("records")
        if not filas_originales:
            return []

        grupos: dict[str, list[dict[str, Any]]] = {}
        sin_incidente: list[dict[str, Any]] = []

        for fila in filas_originales:
            normalizada = self._normalizar_registro(fila)
            incidente = normalizada.get("numero_incidente", "")
            if incidente:
                grupos.setdefault(str(incidente), []).append(fila)
            else:
                sin_incidente.append(fila)

        registros: list[dict[str, Any]] = []
        for incidente, filas_grupo in grupos.items():
            registro = self._normalizar_registro(filas_grupo[0])
            registro["numero_incidente"] = incidente
            registro["importe"] = self._calcular_importe_transacciones(filas_grupo)
            cuenta = self._buscar_primera_cuenta(filas_grupo)
            if cuenta:
                registro["numero_cuenta"] = cuenta
            registros.append(registro)

        registros.extend(self._normalizar_registro(fila) for fila in sin_incidente)
        return registros

    def _buscar_mejor_resolucion(self, caso: dict[str, Any]) -> dict[str, Any] | None:
        """Busca una resolucion entrenada compatible con el caso.

        Args:
            caso: Caso normalizado.

        Returns:
            Resolucion encontrada o `None`.
        """
        canal_caso = str(caso.get("canal", "")).lower()
        motivo_caso = str(caso.get("motivo", "")).lower()

        for resolucion in self.repositorio.listar_resoluciones():
            canal_resolucion = str(resolucion.get("canal", "")).lower()
            motivo_resolucion = str(resolucion.get("motivo", "")).lower()
            canal_match = canal_resolucion in {"", "otros"} or canal_resolucion == canal_caso
            motivo_match = motivo_resolucion in motivo_caso or motivo_caso in motivo_resolucion
            if canal_match and motivo_match:
                return resolucion
        return None

    def _normalizar_registro(self, fila: dict[str, Any]) -> dict[str, Any]:
        """Normaliza nombres de campos frecuentes.

        Args:
            fila: Registro original.

        Returns:
            Registro normalizado para el motor.
        """
        normalizado = {self._normalizar_clave(clave): valor for clave, valor in fila.items()}
        importe = self._obtener_valor(
            normalizado,
            [
                "importe $",
                "importe",
                "monto",
                "monto reclamado",
                "importe reclamado",
            ],
        )
        cuenta = self._obtener_valor(
            normalizado,
            [
                "cuenta desde",
                "cta dde",
                "cta desde",
                "numero_cuenta",
                "nro_cuenta",
                "cuenta",
            ],
        )
        return {
            "numero_incidente": str(
                self._obtener_valor(
                    normalizado,
                    [
                        "numero_incidente",
                        "incidente",
                        "nro_incidente",
                        "nro de operacion",
                        "nro_operacion",
                        "numero_operacion",
                        "numero de operacion",
                        "nro reclamo",
                        "reclamo",
                    ],
                )
                or ""
            ).strip(),
            "nombre_apellido": str(
                self._obtener_valor(
                    normalizado,
                    ["nombre y apellido", "nombre_apellido", "cliente", "pagador", "afectado"],
                )
                or ""
            ).strip(),
            "canal": str(self._obtener_valor(normalizado, ["canal", "tema"]) or "Otros").strip(),
            "motivo": str(
                self._obtener_valor(
                    normalizado,
                    ["motivo", "detalle", "descripcion", "descripcion reclamo"],
                )
                or ""
            ).strip(),
            "fecha": str(
                self._obtener_valor(normalizado, ["fecha", "fecha trx", "fecha operacion"]) or ""
            ).strip(),
            "hora": str(
                self._obtener_valor(normalizado, ["hora", "hora trx", "hora operacion"]) or ""
            ).strip(),
            "importe": parsear_importe(importe),
            "moneda": str(self._obtener_valor(normalizado, ["moneda", "divisa"]) or "Peso").strip(),
            "numero_cuenta": self._normalizar_cuenta(cuenta),
        }

    def convertir_dolar_a_peso(self, importe: float) -> float:
        """Convierte un importe en dolares a pesos segun el coeficiente vigente.

        Args:
            importe: Importe expresado en dolares.

        Returns:
            Importe convertido a pesos.
        """
        return importe * self.coeficiente_dolar

    def _calcular_importe_transacciones(self, filas: list[dict[str, Any]]) -> float:
        """Suma importes marcados; si no hay marca, usa el primer importe disponible.

        Args:
            filas: Filas originales de una misma base o incidente.

        Returns:
            Importe consolidado.
        """
        filas_marcadas = [fila for fila in filas if self._es_transaccion_marcada(fila)]
        filas_a_sumar = filas_marcadas or filas[:1]
        total = 0.0

        for fila in filas_a_sumar:
            normalizada = {self._normalizar_clave(clave): valor for clave, valor in fila.items()}
            total += parsear_importe(
                self._obtener_valor(
                    normalizada,
                    ["importe $", "importe", "monto", "monto reclamado", "importe reclamado"],
                )
            )
        return total

    def _buscar_primera_cuenta(self, filas: list[dict[str, Any]]) -> str:
        """Obtiene la primera cuenta valida dentro de un grupo de filas.

        Args:
            filas: Filas originales.

        Returns:
            Cuenta normalizada a 14 digitos o cadena vacia.
        """
        for fila in filas:
            normalizada = {self._normalizar_clave(clave): valor for clave, valor in fila.items()}
            cuenta = self._normalizar_cuenta(
                self._obtener_valor(
                    normalizada,
                    ["cuenta desde", "cta dde", "cta desde", "numero_cuenta", "nro_cuenta"],
                )
            )
            if cuenta:
                return cuenta
        return ""

    def _es_transaccion_marcada(self, fila: dict[str, Any]) -> bool:
        """Detecta filas seleccionadas o resaltadas en bases exportadas.

        Args:
            fila: Fila original de Excel/CSV.

        Returns:
            `True` si la fila tiene una marca reconocible.
        """
        normalizada = {self._normalizar_clave(clave): valor for clave, valor in fila.items()}
        valor = self._obtener_valor(
            normalizada,
            ["check", "seleccionado", "seleccionada", "resaltado", "resaltada", "marcado"],
        )
        if isinstance(valor, bool):
            return valor
        if valor is None:
            return False
        texto = str(valor).strip().lower()
        return texto in {"1", "x", "si", "sí", "true", "ok", "marcado", "resaltado"}

    def _normalizar_cuenta(self, valor: Any) -> str:
        """Convierte una cuenta a texto de 14 digitos cuando es posible.

        Args:
            valor: Cuenta original.

        Returns:
            Cuenta normalizada o cadena vacia.
        """
        if valor is None:
            return ""
        digitos = re.sub(r"\D", "", str(valor))
        if not digitos:
            return ""
        return digitos[-14:].zfill(14)

    def _normalizar_clave(self, clave: Any) -> str:
        """Normaliza encabezados de CSV/Excel para matchear variantes.

        Args:
            clave: Encabezado original.

        Returns:
            Encabezado comparable.
        """
        texto = unicodedata.normalize("NFKD", str(clave))
        texto = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
        texto = re.sub(r"[^a-zA-Z0-9$]+", " ", texto).strip().lower()
        return re.sub(r"\s+", " ", texto)

    def _actualizar_coeficiente_dolar(self, dataframe: pd.DataFrame) -> None:
        """Busca y actualiza el coeficiente dolar desde RESOLUCIONES APC.

        Args:
            dataframe: Hoja del Excel de resoluciones.
        """
        if dataframe.empty:
            return

        columnas = {self._normalizar_clave(columna): columna for columna in dataframe.columns}
        columna_coeficiente = self._buscar_columna(
            columnas,
            ["coeficiente dolar", "coeficiente_dolar", "coef dolar", "dolar"],
        )
        if columna_coeficiente is None:
            return

        for valor in dataframe[columna_coeficiente].dropna():
            coeficiente = parsear_importe(valor)
            if coeficiente > 0:
                self.coeficiente_dolar = coeficiente
                self.logger.info("Coeficiente dolar actualizado desde Excel de resoluciones")
                return

    def _buscar_columna(self, columnas: dict[str, Any], aliases: list[str]) -> Any:
        """Busca una columna por aliases normalizados.

        Args:
            columnas: Mapa de columnas normalizadas a columnas originales.
            aliases: Nombres posibles.

        Returns:
            Columna original encontrada o `None`.
        """
        for alias in aliases:
            columna = columnas.get(self._normalizar_clave(alias))
            if columna is not None:
                return columna
        return None

    def _obtener_valor(self, normalizado: dict[str, Any], aliases: list[str]) -> Any:
        """Busca el primer valor no vacio entre aliases de columna.

        Args:
            normalizado: Diccionario con claves ya normalizadas.
            aliases: Nombres posibles de columna.

        Returns:
            Primer valor encontrado o `None`.
        """
        for alias in aliases:
            valor = normalizado.get(self._normalizar_clave(alias))
            if valor is not None and str(valor).strip().lower() != "nan":
                return valor
        return None

    def _inferir_canal_desde_tema(self, tema: str) -> str:
        """Infiere canal APC a partir del tema de una resolucion.

        Args:
            tema: Tema textual de la resolucion.

        Returns:
            Canal normalizado.
        """
        tema_upper = tema.upper()
        if "ATM" in tema_upper:
            return "ATM"
        if "MPOS" in tema_upper:
            return "mPOS"
        if "POS" in tema_upper:
            return "POS"
        if "BNA" in tema_upper:
            return "BNA+"
        if "MODO" in tema_upper:
            return "MODO"
        if "CASH" in tema_upper:
            return "Cash In"
        if "ECOMMERCE" in tema_upper or "LINKGO" in tema_upper:
            return "eCommerce"
        return "Otros"

    def _generar_diario_desde_nota(self, tema: str, nota: str) -> str:
        """Genera una base editable para Diario a partir de una nota.

        Args:
            tema: Tema de resolucion.
            nota: Nota de resolucion.

        Returns:
            Texto Diario sugerido.
        """
        return f"Se analiza caso APC bajo tema '{tema}'. Resolucion sugerida: {nota}"
