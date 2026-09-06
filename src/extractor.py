"""Extraccion de datos desde documentos de reclamos APC."""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

from src.utils import buscar_primer_match, normalizar_texto, parsear_importe


class Extractor:
    """Extractor de texto y datos criticos desde archivos locales."""

    def __init__(self, logger: logging.Logger | None = None) -> None:
        """Inicializa el extractor.

        Args:
            logger: Logger opcional para registrar acciones y errores.
        """
        self.logger = logger or logging.getLogger(self.__class__.__name__)
        self._contenido_reciente: str = ""

    def leer_pdf(self, archivo: str) -> str:
        """Lee texto desde un archivo PDF.

        Args:
            archivo: Ruta del archivo PDF.

        Returns:
            Texto extraido del PDF.

        Raises:
            RuntimeError: Si `pypdf` no esta instalado o el PDF no puede leerse.
        """
        path = Path(archivo)
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            texto = "\n".join(page.extract_text() or "" for page in reader.pages)
            self._contenido_reciente = normalizar_texto(texto)
            self.logger.info("PDF leido: %s", path)
            return self._contenido_reciente
        except Exception as exc:
            self.logger.exception("No se pudo leer el PDF %s", path)
            raise RuntimeError(f"No se pudo leer el PDF {path}: {exc}") from exc

    def leer_excel(self, archivo: str) -> dict[str, list[dict[str, Any]]]:
        """Lee un archivo Excel y devuelve filas por hoja.

        Args:
            archivo: Ruta del archivo Excel.

        Returns:
            Diccionario con nombre de hoja y filas como diccionarios.

        Raises:
            RuntimeError: Si `openpyxl` no esta instalado o el archivo falla.
        """
        path = Path(archivo)
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(path, data_only=True)
            resultado: dict[str, list[dict[str, Any]]] = {}

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                filas = list(sheet.iter_rows(values_only=True))
                if not filas:
                    resultado[sheet_name] = []
                    continue

                encabezados = [str(valor or "").strip() for valor in filas[0]]
                registros: list[dict[str, Any]] = []
                for fila in filas[1:]:
                    registros.append(
                        {
                            encabezados[indice] or f"columna_{indice + 1}": valor
                            for indice, valor in enumerate(fila)
                        }
                    )
                resultado[sheet_name] = registros

            self._contenido_reciente = normalizar_texto(str(resultado))
            self.logger.info("Excel leido: %s", path)
            return resultado
        except Exception as exc:
            self.logger.exception("No se pudo leer el Excel %s", path)
            raise RuntimeError(f"No se pudo leer el Excel {path}: {exc}") from exc

    def leer_txt(self, archivo: str) -> str:
        """Lee texto desde un archivo TXT.

        Args:
            archivo: Ruta del archivo TXT.

        Returns:
            Contenido normalizado del archivo.
        """
        path = Path(archivo)
        try:
            texto = path.read_text(encoding="utf-8")
            self._contenido_reciente = normalizar_texto(texto)
            self.logger.info("TXT leido: %s", path)
            return self._contenido_reciente
        except Exception as exc:
            self.logger.exception("No se pudo leer el TXT %s", path)
            raise RuntimeError(f"No se pudo leer el TXT {path}: {exc}") from exc

    def extraer_fecha(self) -> str:
        """Extrae una fecha del ultimo contenido leido.

        Returns:
            Fecha encontrada en formato textual.
        """
        return buscar_primer_match(
            [
                r"fecha[:\s]+(\d{4}-\d{2}-\d{2})",
                r"fecha[:\s]+(\d{2}/\d{2}/\d{4})",
                r"\b(\d{4}-\d{2}-\d{2})\b",
                r"\b(\d{2}/\d{2}/\d{4})\b",
            ],
            self._contenido_reciente,
        )

    def extraer_importe(self) -> float:
        """Extrae un importe del ultimo contenido leido.

        Returns:
            Importe como `float`. Devuelve `0.0` si no encuentra datos.
        """
        valor = buscar_primer_match(
            [
                r"importe[:\s$ARS]+([\d.,]+)",
                r"monto[:\s$ARS]+([\d.,]+)",
                r"\$[\s]*([\d.,]+)",
            ],
            self._contenido_reciente,
        )
        return parsear_importe(valor)

    def extraer_numero_cuenta(self) -> str:
        """Extrae un numero de cuenta del ultimo contenido leido.

        Returns:
            Numero de cuenta como texto.
        """
        return buscar_primer_match(
            [
                r"numero de cuenta[:\s]+([A-Z0-9-]{6,})",
                r"cuenta[:\s]+([A-Z0-9-]{6,})",
                r"\b(\d{10,22})\b",
            ],
            self._contenido_reciente,
        )

    def extraer_hora(self) -> str:
        """Extrae una hora del ultimo contenido leido.

        Returns:
            Hora en formato HH:MM si existe.
        """
        match = re.search(r"\b([01]\d|2[0-3]):([0-5]\d)\b", self._contenido_reciente)
        return match.group(0) if match else ""

    def extraer_motivo(self) -> str:
        """Extrae el motivo declarado en el ultimo contenido leido.

        Returns:
            Motivo del reclamo o cadena vacia.
        """
        return buscar_primer_match(
            [
                r"motivo[:\s]+(.+?)(?:\.|$)",
                r"causa[:\s]+(.+?)(?:\.|$)",
            ],
            self._contenido_reciente,
        )
